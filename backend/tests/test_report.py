"""Tests for GET /api/v1/report (Phase 12) — auth, both formats' file
signatures, filter behavior (verified by reading the generated .xlsx back
with openpyxl), and the empty-history path for a fresh user."""

import io

from openpyxl import load_workbook

PREDICT_PAYLOAD = {
    "crop_type": "wheat",
    "soil_moisture_pct": 22.0,
    "canal_flow_cusecs": 300.0,
    "use_live_weather": False,
    "manual_temperature_c": 32.0,
    "manual_humidity_pct": 40.0,
    "manual_rainfall_mm": 1.0,
    "manual_evapotranspiration_mm": 6.0,
}


def _predict(client, headers, district: str):
    response = client.post("/api/v1/predict", json={**PREDICT_PAYLOAD, "district": district}, headers=headers)
    assert response.status_code == 200, response.text
    return response.json()


def test_report_requires_auth(client):
    response = client.get("/api/v1/report")

    assert response.status_code == 401


def test_xlsx_report_is_a_real_workbook(client, register_user):
    token = register_user(username="repxlsx")
    headers = {"Authorization": f"Bearer {token}"}
    _predict(client, headers, "Lahore")

    response = client.get("/api/v1/report", headers=headers)

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert response.content[:2] == b"PK"  # ZIP magic bytes — .xlsx is a ZIP container
    assert len(response.content) > 0
    assert "attachment" in response.headers["content-disposition"]
    assert ".xlsx" in response.headers["content-disposition"]

    # A real, readable workbook with the expected sheets — not just bytes
    # that happen to start with "PK".
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Summary", "Weather & Forecast", "History"]


def test_pdf_report_starts_with_pdf_magic_bytes(client, register_user):
    token = register_user(username="reppdf")
    headers = {"Authorization": f"Bearer {token}"}
    _predict(client, headers, "Lahore")

    response = client.get("/api/v1/report", headers=headers, params={"format": "pdf"})

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content[:4] == b"%PDF"
    assert len(response.content) > 0
    assert ".pdf" in response.headers["content-disposition"]


def test_report_rejects_unknown_format(client, register_user):
    token = register_user(username="repbadformat")
    headers = {"Authorization": f"Bearer {token}"}

    response = client.get("/api/v1/report", headers=headers, params={"format": "docx"})

    assert response.status_code == 422


def test_empty_history_report_still_returns_200(client, register_user):
    """A fresh user with zero predictions must still get a real report
    (weather/forecast sections + an explicit 'no predictions' row) —
    never an error."""
    token = register_user(username="repfresh")
    headers = {"Authorization": f"Bearer {token}"}

    xlsx_response = client.get("/api/v1/report", headers=headers)
    assert xlsx_response.status_code == 200
    wb = load_workbook(io.BytesIO(xlsx_response.content))
    history_ws = wb["History"]
    assert history_ws.cell(row=2, column=1).value == "No predictions recorded yet."

    pdf_response = client.get("/api/v1/report", headers=headers, params={"format": "pdf"})
    assert pdf_response.status_code == 200
    assert pdf_response.content[:4] == b"%PDF"


def test_district_filter_narrows_history_table(client, register_user):
    """Seeds 3 predictions across 2 districts, then asserts the report's
    History sheet contains only the filtered district's rows — read back
    with openpyxl, not just trusting the endpoint returned 200."""
    token = register_user(username="repfilter")
    headers = {"Authorization": f"Bearer {token}"}
    _predict(client, headers, "Lahore")
    _predict(client, headers, "Lahore")
    _predict(client, headers, "Multan")

    response = client.get("/api/v1/report", headers=headers, params={"district": "Lahore"})
    assert response.status_code == 200

    wb = load_workbook(io.BytesIO(response.content))
    history_ws = wb["History"]
    data_rows = [row for row in history_ws.iter_rows(min_row=2, values_only=True) if row[0] is not None]
    assert len(data_rows) == 2
    assert all(row[1] == "Lahore" for row in data_rows)


def test_field_filter_scopes_report_to_one_field(client, register_user):
    token = register_user(username="repfieldfilter")
    headers = {"Authorization": f"Bearer {token}"}

    field_response = client.post(
        "/api/v1/fields",
        json={"name": "North Plot", "district": "Multan", "crop_type": "wheat"},
        headers=headers,
    )
    assert field_response.status_code == 201, field_response.text
    field_id = field_response.json()["id"]

    predict_response = client.post(
        "/api/v1/predict",
        json={**PREDICT_PAYLOAD, "district": "Multan", "field_id": field_id},
        headers=headers,
    )
    assert predict_response.status_code == 200, predict_response.text
    # An unrelated prediction with no field_id must not appear in the filtered report.
    _predict(client, headers, "Multan")

    response = client.get("/api/v1/report", headers=headers, params={"field_id": field_id})
    assert response.status_code == 200

    wb = load_workbook(io.BytesIO(response.content))
    history_ws = wb["History"]
    data_rows = [row for row in history_ws.iter_rows(min_row=2, values_only=True) if row[0] is not None]
    assert len(data_rows) == 1

    summary_ws = wb["Summary"]
    field_names = [row[1] for row in summary_ws.iter_rows(values_only=True) if row[0] == "Name"]
    assert "North Plot" in field_names


def test_report_404s_for_another_users_field_id(client, register_user):
    token_a = register_user(username="repownerA")
    token_b = register_user(username="repownerB")

    field_response = client.post(
        "/api/v1/fields",
        json={"name": "A's Plot", "district": "Lahore", "crop_type": "wheat"},
        headers={"Authorization": f"Bearer {token_a}"},
    )
    field_id = field_response.json()["id"]

    response = client.get(
        "/api/v1/report", headers={"Authorization": f"Bearer {token_b}"}, params={"field_id": field_id}
    )

    assert response.status_code == 404


def test_report_rejects_unknown_district(client, register_user):
    token = register_user(username="repbaddistrict")
    headers = {"Authorization": f"Bearer {token}"}

    response = client.get("/api/v1/report", headers=headers, params={"district": "Atlantis"})

    assert response.status_code == 400
